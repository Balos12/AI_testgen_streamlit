import streamlit as st
import openai
import os
from docx import Document
from dotenv import load_dotenv
from io import BytesIO

load_dotenv()
openai.api_key = os.getenv("OPENAI_API_KEY")

# Многоязычность без po
LANGUAGES = {
    "en": {
        "title": "Programming Test",
        "enter_topic": "Enter the test topic:",
        "select_difficulty": "Select difficulty level:",
        "number_of_questions": "Number of questions",
        "generate_test": "Generate test",
        "take_the_test": "📋 Take the test:",
        "finish_test": "Finish test",
        "result": "Result",
        "correct": "Correct",
        "incorrect": "Your answer: {} Correct answer: {}",
        "download_docx": "Download DOCX",
        "choose_answer": "Choose an answer:",
        "difficulty_levels": ["Easy", "Medium", "Hard"],  # added for difficulty
        'explanations':"explanations"
    },
    "ru": {
        "title": "Программируемый тест",
        "enter_topic": "Введите тему теста:",
        "select_difficulty": "Выберите уровень сложности:",
        "number_of_questions": "Количество вопросов",
        "generate_test": "Сгенерировать тест",
        "take_the_test": "📋 Пройдите тест:",
        "finish_test": "Завершить тест",
        "result": "Результат",
        "correct": "Правильный",
        "incorrect": "Ваш ответ: {} Правильный ответ: {}",
        "download_docx": "Скачать DOCX",
        "choose_answer": "Выберите ответ:",
        "difficulty_levels": ["Лёгкий", "Средний", "Трудный"],  # added for difficulty
        'explanations':"объяснения"
    }
}

# Выбор языка
language = st.selectbox("Choose Language", ["en", "ru"])

# Интерфейс
st.title(LANGUAGES[language]["title"])

# Ввод темы и уровня сложности
topic = st.text_input(LANGUAGES[language]["enter_topic"])
difficulty = st.selectbox(LANGUAGES[language]["select_difficulty"], LANGUAGES[language]["difficulty_levels"])
num_questions = st.slider(LANGUAGES[language]["number_of_questions"], 1, 10, 3)

if 'test_data' not in st.session_state:
    st.session_state.test_data = []

if st.button(LANGUAGES[language]["generate_test"]):
    if topic:
        with st.spinner(LANGUAGES[language]["generate_test"]):
            prompt = f"""
You are a programming test generator. Generate {num_questions} questions on the topic "{topic}".
Difficulty: {difficulty}. Format:
Question 1: ...
A) ...
B) ...
C) ...
D) ...
Answer: X
"""

            response = openai.ChatCompletion.create(
                model="gpt-4o",
                messages=[{"role": "user", "content": prompt}],
                temperature=0.7,
                max_tokens=2000
            )

            test_text = response['choices'][0]['message']['content']
            st.session_state.test_data = test_text.split("\n\n")
            st.success(LANGUAGES[language]["generate_test"])

# Прохождение теста
if st.session_state.test_data:
    st.subheader(LANGUAGES[language]["take_the_test"])

    user_answers = {}
    correct_answers = {}
    questions = []

    for i, block in enumerate(st.session_state.test_data):
        if "Question" in block:
            lines = block.split("\n")
            question = lines[0]
            options = lines[1:5]
            answer_line = [l for l in lines if "Answer:" in l][0]
            correct = answer_line.split(": ")[1].strip()

            questions.append(question)
            correct_answers[question] = correct

            st.write(question)
            for option in options:
                st.write(option)

            user_answer = st.radio(LANGUAGES[language]["choose_answer"], ("A", "B", "C", "D"), key=f"q{i}")
            user_answers[question] = user_answer
            st.write("---")

    if st.button(LANGUAGES[language]["finish_test"]):
        score = 0
        total = len(questions)
        result_output = ""

        for q in questions:
            user_ans = user_answers[q]
            correct_ans = correct_answers[q]

            if user_ans == correct_ans:
                score += 1
                result_output += f"✅ {q} — {LANGUAGES[language]['correct']} ({user_ans})\n"
            else:
                result_output += f"❌ {q} — {LANGUAGES[language]['incorrect'].format(user_ans, correct_ans)}\n"

        percent = (score / total) * 100

        st.subheader(f"{LANGUAGES[language]['result']}: {score} / {total} ({percent:.2f}%)")
        st.text(result_output)

        # Генерация пояснений одним запросом
        st.subheader("📝 Explanations:")

        explain_prompt = f"""Give simple "{topic}" explanations for the correct answers to the following programming test questions:\n\n"""
        for q in questions:
            correct_ans = correct_answers[q]
            explain_prompt += f"Question: {q}\nCorrect answer: {correct_ans}\n\n"

        explain_response = openai.ChatCompletion.create(
            model="gpt-4o",
            messages=[{"role": "user", "content": explain_prompt}],
            temperature=0.5,
            max_tokens=3000
        )

        explanations_text = explain_response['choices'][0]['message']['content']

        # Показ пояснений как единый текст
        st.markdown(explanations_text)

        # Экспорт в DOCX
        def export_to_docx(test_blocks, explanations_text):
            doc = Document()

            # Apply DejaVu Sans font to the entire document
            style = doc.styles['Normal']
            font = style.font
            font.name = 'DejaVu Sans'  # Apply DejaVu Sans font

            doc.add_heading("Programming Test", 0)

            for block in test_blocks:
                doc.add_paragraph(block)
                doc.add_paragraph("")  # Add space between questions

            # Add explanations
            doc.add_paragraph("📝 Explanations:")
            doc.add_paragraph(explanations_text)

            buffer = BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            return buffer

        # Генерация и экспорт
        docx_file = export_to_docx(st.session_state.test_data, explanations_text)
        st.download_button(
            label="📄 " + LANGUAGES[language]["download_docx"],
            data=docx_file,
            file_name="test_with_explanations.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
